from __future__ import annotations

import csv
import math
import re
import threading
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import load_workbook


class AircraftType(str, Enum):
    CFM = "CFM"
    LEAP = "LEAP"
    PW = "PW"


class AircraftCategory(str, Enum):
    A319 = "A319"
    A320_A321 = "A320、A321"
    A321NEO = "A321NEO"


SUPPORTED_AIRCRAFT: dict[str, tuple[AircraftType, AircraftCategory]] = {
    "A319CFM": (AircraftType.CFM, AircraftCategory.A319),
    "A319LEAP": (AircraftType.LEAP, AircraftCategory.A319),
    "A320CFM": (AircraftType.CFM, AircraftCategory.A320_A321),
    "A321CFM": (AircraftType.CFM, AircraftCategory.A320_A321),
    "A320PW": (AircraftType.PW, AircraftCategory.A320_A321),
    "A320LEAP": (AircraftType.LEAP, AircraftCategory.A320_A321),
    "A321LEAP": (AircraftType.LEAP, AircraftCategory.A321NEO),
    "A321PW": (AircraftType.PW, AircraftCategory.A321NEO),
}


class AnalysisCancelled(Exception):
    """Raised internally when the user requests a safe cancellation."""


@dataclass(frozen=True)
class AircraftConfig:
    marker: str
    altitude_col: int
    stick_cols: tuple[int, int]
    rudder_cols: tuple[int, int]
    aircraft_heading_col: int
    runway_heading_col: int
    takeoff_group_1: tuple[int, int]
    takeoff_group_2: tuple[int, int]
    takeoff_min_1: int
    takeoff_min_2: int
    pitch_rate_cols: tuple[int, int]
    baro_altitude_col: int
    wind_speed_col: int
    wind_direction_col: int


@dataclass
class MetricResult:
    success: bool
    value: Any = None
    error_code: str = ""
    error_message: str = ""
    applicable: bool = True


@dataclass
class FlightResult:
    person: str
    path: Path
    aircraft_type: AircraftType
    flight_time: datetime | None
    flight_number: str | None
    metrics: dict[str, MetricResult]
    crosswind_metric: MetricResult
    touchdown_row: int | None = None
    fifty_ft_row: int | None = None
    takeoff_row: int | None = None
    aircraft_heading: float | None = None
    angle_difference: float | None = None
    signed_angle: float | None = None
    reverse_rudder_value: float | None = None
    reverse_rudder_row: int | None = None
    max_pitch_rate: float | None = None
    baro_altitude: float | None = None
    runway_heading: float | None = None
    wind_speed: float | None = None
    wind_direction: float | None = None
    crosswind: float | None = None
    wind_source_row: int | None = None
    wind_row_offset: int | None = None
    warnings: list[str] = field(default_factory=list)
    registration_raw: str = ""
    registration: str = ""
    database_aircraft_type: str = ""
    aircraft_category: str = ""
    validation_status: str = "VALID"


@dataclass
class MetricStats:
    candidate: int = 0
    valid: int = 0
    triggered: int = 0
    failed: int = 0
    skipped: int = 0


@dataclass
class PersonSummary:
    person: str
    aircraft_category: str
    stick_exceed_count: int
    stick_exceed_ratio: float | None
    rudder_exceed_count: int
    rudder_valid_sample_count: int
    rudder_exceed_ratio: float | None
    cross_angle_count: int
    cross_angle_ratio: float | None
    pitch_rate_count: int
    pitch_rate_ratio: float | None
    high_altitude_count: int
    high_altitude_ratio: float | None
    reverse_rudder_count: int
    stats: dict[str, MetricStats]


@dataclass
class LogRecord:
    record_type: str
    timestamp: datetime
    person: str = ""
    file_path: str = ""
    aircraft_type: str = ""
    metric: str = ""
    status: str = ""
    error_code: str = ""
    message: str = ""
    candidate: str = ""
    valid: str = ""
    triggered: str = ""
    failed: str = ""
    excel_file: str = ""
    docx_status: str = ""
    wind_source_row: str = ""
    wind_row_offset: str = ""
    crosswind: str = ""
    signed_angle: str = ""
    cross_direction: str = ""
    reverse_rudder_value: str = ""
    reverse_rudder_row: str = ""
    fifty_ft_row: str = ""
    touchdown_row: str = ""
    reverse_rudder_triggered: str = ""
    aircraft_database: str = ""
    registration_raw: str = ""
    registration: str = ""
    database_aircraft_type: str = ""
    aircraft_category: str = ""
    filename_engine_marker: str = ""
    skipped: str = ""


@dataclass
class ProgressEvent:
    stage: str
    person: str = ""
    file_name: str = ""
    completed: int = 0
    total: int = 0
    message: str = ""


@dataclass
class AnalysisResult:
    selected_folder: Path
    aircraft_database: Path
    started_at: datetime
    ended_at: datetime
    summaries: list[PersonSummary]
    flights: list[FlightResult]
    log_records: list[LogRecord]
    cancelled: bool = False
    preflight_failed: bool = False


@dataclass
class AircraftDatabaseLoadResult:
    success: bool
    path: Path
    registrations: dict[str, str] = field(default_factory=dict)
    error_code: str = ""
    error_message: str = ""


@dataclass
class PreflightFileResult:
    person: str
    path: Path
    aircraft_type: AircraftType
    registration_raw: str = ""
    registration: str = ""
    database_aircraft_type: str = ""
    aircraft_category: str = ""
    status: str = "VALID"
    error_code: str = ""
    message: str = ""


@dataclass
class PreflightResult:
    success: bool
    files: list[PreflightFileResult]
    blocking_files: list[PreflightFileResult] = field(default_factory=list)


def _column_number(label: str) -> int:
    value = 0
    for char in label.upper():
        value = value * 26 + ord(char) - ord("A") + 1
    return value - 1


CONFIGS: dict[AircraftType, AircraftConfig] = {
    AircraftType.CFM: AircraftConfig(
        marker="YZDCFM",
        altitude_col=_column_number("D"),
        stick_cols=(_column_number("AM"), _column_number("BB")),
        rudder_cols=(_column_number("BP"), _column_number("BW")),
        aircraft_heading_col=_column_number("BG"),
        runway_heading_col=_column_number("BH"),
        takeoff_group_1=(_column_number("D"), _column_number("J")),
        takeoff_group_2=(_column_number("K"), _column_number("Q")),
        takeoff_min_1=2,
        takeoff_min_2=2,
        pitch_rate_cols=(_column_number("BL"), _column_number("BO")),
        baro_altitude_col=_column_number("BK"),
        wind_speed_col=_column_number("BI"),
        wind_direction_col=_column_number("BJ"),
    ),
    AircraftType.LEAP: AircraftConfig(
        marker="YZDLEAP",
        altitude_col=_column_number("D"),
        stick_cols=(_column_number("AE"), _column_number("AT")),
        rudder_cols=(_column_number("BF"), _column_number("BM")),
        aircraft_heading_col=_column_number("AW"),
        runway_heading_col=_column_number("AX"),
        takeoff_group_1=(_column_number("D"), _column_number("H")),
        takeoff_group_2=(_column_number("I"), _column_number("M")),
        takeoff_min_1=1,
        takeoff_min_2=1,
        pitch_rate_cols=(_column_number("BB"), _column_number("BE")),
        baro_altitude_col=_column_number("BA"),
        wind_speed_col=_column_number("AY"),
        wind_direction_col=_column_number("AZ"),
    ),
    AircraftType.PW: AircraftConfig(
        marker="YZDPW",
        altitude_col=_column_number("D"),
        stick_cols=(_column_number("AE"), _column_number("AT")),
        rudder_cols=(_column_number("BH"), _column_number("BO")),
        aircraft_heading_col=_column_number("AY"),
        runway_heading_col=_column_number("AZ"),
        takeoff_group_1=(_column_number("D"), _column_number("H")),
        takeoff_group_2=(_column_number("I"), _column_number("M")),
        takeoff_min_1=1,
        takeoff_min_2=2,
        pitch_rate_cols=(_column_number("BD"), _column_number("BG")),
        baro_altitude_col=_column_number("BC"),
        wind_speed_col=_column_number("BA"),
        wind_direction_col=_column_number("BB"),
    ),
}

METRIC_LABELS = {
    "stick": "300ft以下杆量",
    "rudder": "50ft以下舵量",
    "cross_angle": "接地交叉角",
    "pitch_rate": "PF抬头速率",
    "high_altitude": "高高原抬头速率",
    "reverse_rudder": "疑似接地前反向蹬舵",
}

_TIME_PATTERN = re.compile(r"(?<!\d)(\d{14})(?!\d)")
_FLIGHT_NUMBER_PATTERN = re.compile(r"CA\d+", re.IGNORECASE)
_REGISTRATION_PATTERN = re.compile(r"B-[A-Z0-9]{4}", re.IGNORECASE)


def parse_flight_time(name: str) -> datetime | None:
    for match in _TIME_PATTERN.finditer(name):
        try:
            return datetime.strptime(match.group(1), "%Y%m%d%H%M%S")
        except ValueError:
            continue
    return None


def parse_flight_number(name: str, aircraft_type: AircraftType | None = None) -> str | None:
    upper_name = name.upper()
    if aircraft_type is not None:
        marker_index = upper_name.find(CONFIGS[aircraft_type].marker)
        if marker_index >= 0:
            match = _FLIGHT_NUMBER_PATTERN.search(name, marker_index + len(CONFIGS[aircraft_type].marker))
            if match:
                return match.group(0).upper()
    match = _FLIGHT_NUMBER_PATTERN.search(name)
    return match.group(0).upper() if match else None


def identify_aircraft(name: str) -> AircraftType | None:
    upper_name = name.upper()
    matches = [kind for kind, config in CONFIGS.items() if config.marker in upper_name]
    return matches[0] if len(matches) == 1 else None


def parse_registration(name: str) -> tuple[str, str] | None:
    match = _REGISTRATION_PATTERN.search(name)
    if match is None:
        return None
    raw = match.group(0)
    return raw, raw.strip().upper()


def load_aircraft_database(path: str | Path) -> AircraftDatabaseLoadResult:
    database_path = Path(path).resolve()
    if database_path.suffix.lower() != ".xlsx":
        return AircraftDatabaseLoadResult(False, database_path, error_code="DATABASE_EXTENSION", error_message="飞机数据库必须是.xlsx文件")
    try:
        workbook = load_workbook(database_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError("数据库为空")
        normalized_headers = [str(value).strip() if value is not None else "" for value in header]
        try:
            registration_col = normalized_headers.index("Registration")
            aircraft_col = normalized_headers.index("A/C TYPE")
        except ValueError as exc:
            raise ValueError("数据库必须包含 Registration 和 A/C TYPE 列") from exc
        registrations: dict[str, str] = {}
        conflicts: set[str] = set()
        for row in rows:
            raw_registration = row[registration_col] if registration_col < len(row) else None
            raw_type = row[aircraft_col] if aircraft_col < len(row) else None
            if raw_registration is None or not str(raw_registration).strip():
                continue
            registration = str(raw_registration).strip().upper()
            aircraft_type = "" if raw_type is None else str(raw_type).strip().upper()
            previous = registrations.get(registration)
            if previous is not None and previous != aircraft_type:
                conflicts.add(registration)
            registrations[registration] = aircraft_type
        workbook.close()
        if conflicts:
            names = "、".join(sorted(conflicts))
            raise ValueError(f"重复注册号对应不同机型：{names}")
        return AircraftDatabaseLoadResult(True, database_path, registrations)
    except Exception as exc:
        return AircraftDatabaseLoadResult(False, database_path, error_code="DATABASE_READ_ERROR", error_message=str(exc))


def preflight_candidates(
    people: list[tuple[Path, list[tuple[Path, AircraftType]]]],
    database: AircraftDatabaseLoadResult,
) -> PreflightResult:
    results: list[PreflightFileResult] = []
    blocking: list[PreflightFileResult] = []
    if not database.success:
        return PreflightResult(False, results)
    for person_dir, files in people:
        for path, filename_engine in files:
            parsed = parse_registration(path.name)
            if parsed is None:
                item = PreflightFileResult(person_dir.name, path, filename_engine, status="BLOCKED", error_code="REGISTRATION_MISSING", message="文件名中缺少符合B-[A-Z0-9]{4}的注册号")
                results.append(item)
                blocking.append(item)
                continue
            raw, registration = parsed
            database_type = database.registrations.get(registration)
            if database_type is None:
                item = PreflightFileResult(person_dir.name, path, filename_engine, raw, registration, status="BLOCKED", error_code="REGISTRATION_NOT_FOUND", message="注册号未在飞机数据库中找到")
                results.append(item)
                blocking.append(item)
                continue
            mapping = SUPPORTED_AIRCRAFT.get(database_type)
            if mapping is None:
                results.append(PreflightFileResult(person_dir.name, path, filename_engine, raw, registration, database_type, status="SKIPPED", error_code="AIRCRAFT_TYPE_UNSUPPORTED", message="数据库具体机型不属于支持的8种机型"))
                continue
            expected_engine, category = mapping
            if expected_engine != filename_engine:
                results.append(PreflightFileResult(person_dir.name, path, filename_engine, raw, registration, database_type, category.value, "SKIPPED", "ENGINE_MARKER_CONFLICT", f"数据库发动机={expected_engine.value}，文件名发动机={filename_engine.value}"))
                continue
            results.append(PreflightFileResult(person_dir.name, path, filename_engine, raw, registration, database_type, category.value))
    return PreflightResult(not blocking, results, blocking)


def discover_people(selected_folder: Path) -> list[tuple[Path, list[tuple[Path, AircraftType]]]]:
    if not selected_folder.exists() or not selected_folder.is_dir():
        raise ValueError(f"所选数据目录不存在或不是文件夹：{selected_folder}")

    people: list[tuple[Path, list[tuple[Path, AircraftType]]]] = []
    person_dirs = sorted(
        (path for path in selected_folder.iterdir() if path.is_dir()),
        key=lambda path: (path.name.casefold(), path.name),
    )
    for person_dir in person_dirs:
        flights: list[tuple[Path, AircraftType]] = []
        for path in person_dir.iterdir():
            if not path.is_file() or path.suffix.lower() != ".csv":
                continue
            aircraft_type = identify_aircraft(path.name)
            if aircraft_type is not None:
                flights.append((path, aircraft_type))
        flights.sort(
            key=lambda item: (
                parse_flight_time(item[0].name) is None,
                parse_flight_time(item[0].name) or datetime.max,
                item[0].name.casefold(),
                item[0].name,
            )
        )
        people.append((person_dir, flights))
    return people


def _finite_number(value: str | None) -> float | None:
    if value is None:
        return None
    try:
        number = float(value.strip())
    except (AttributeError, TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _cell(row: list[str], index: int) -> str | None:
    return row[index] if 0 <= index < len(row) else None


def _numbers(row: list[str], start: int, end: int) -> list[float | None]:
    return [_finite_number(_cell(row, index)) for index in range(start, end + 1)]


def _select_encoding(path: Path, cancel_event: threading.Event) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp936"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                while handle.read(1024 * 1024):
                    if cancel_event.is_set():
                        raise AnalysisCancelled
            return encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeError("文件无法使用支持的编码读取")


def _failed_metrics(code: str, message: str) -> dict[str, MetricResult]:
    return {
        "stick": MetricResult(False, error_code=code, error_message=message),
        "rudder": MetricResult(False, error_code=code, error_message=message),
        "cross_angle": MetricResult(False, error_code=code, error_message=message),
        "reverse_rudder": MetricResult(False, error_code=code, error_message=message),
        "pitch_rate": MetricResult(False, error_code=code, error_message=message),
        "high_altitude": MetricResult(
            True,
            value=None,
            error_code="NOT_APPLICABLE",
            error_message="抬头速率是否超限无法确定",
            applicable=False,
        ),
    }


def calculate_crosswind(wind_speed: float, wind_direction: float, runway_heading: float) -> float:
    """Calculate crosswind in kt; negative is from the left and positive is from the right."""
    normalized_direction = wind_direction % 360
    normalized_runway = runway_heading % 360
    value = wind_speed * math.sin(math.radians(normalized_direction - normalized_runway))
    return 0.0 if abs(value) < 1e-12 else value


def analyze_flight(
    person: str,
    path: Path,
    aircraft_type: AircraftType,
    cancel_event: threading.Event,
) -> FlightResult:
    config = CONFIGS[aircraft_type]
    flight_time = parse_flight_time(path.name)
    flight_number = parse_flight_number(path.name, aircraft_type)
    warnings: list[str] = []
    expected_width: int | None = None
    short_rows = 0
    long_rows = 0

    seen_above_2000 = False
    approach_rows: list[
        tuple[
            int,
            int,
            float | None,
            list[float | None],
            list[float | None],
            float | None,
            float | None,
        ]
    ] = []
    rudder_columns_available = False
    touchdown_data_index: int | None = None
    touchdown_csv_row: int | None = None
    touchdown_aircraft_heading: float | None = None
    touchdown_runway_heading: float | None = None
    last_runway_heading: float | None = None
    runway_heading_values: set[float] = set()
    last_wind_at_or_before: tuple[int, int, float, float] | None = None
    first_wind_after: tuple[int, int, float, float] | None = None
    landing_needs_wind = False

    takeoff_data_index: int | None = None
    takeoff_csv_row: int | None = None
    previous_pitch_rates: list[float | None] | None = None
    pitch_window: list[float | None] = []
    pitch_rows_remaining = 0
    takeoff_baro: float | None = None

    try:
        encoding = _select_encoding(path, cancel_event)
        with path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle)
            try:
                first_row = next(reader)
            except StopIteration:
                raise ValueError("CSV为空文件")
            data_start_csv_row = 3
            if (
                first_row
                and first_row[0].strip().lower().startswith("dataframe info")
            ):
                try:
                    header = next(reader)
                except StopIteration:
                    raise ValueError("CSV包含元数据行，但缺少字段名行")
                data_start_csv_row = 4
            else:
                header = first_row
            try:
                next(reader)  # units row
            except StopIteration:
                raise ValueError("CSV缺少单位行和数据行")

            expected_width = len(header)
            if expected_width == 0:
                raise ValueError("CSV表头为空")
            rudder_columns_available = expected_width > config.rudder_cols[1]

            for data_index, row in enumerate(reader):
                csv_row = data_index + data_start_csv_row
                if data_index % 250 == 0 and cancel_event.is_set():
                    raise AnalysisCancelled

                if len(row) < expected_width:
                    short_rows += 1
                elif len(row) > expected_width:
                    long_rows += 1

                runway_value = _finite_number(_cell(row, config.runway_heading_col))
                if runway_value is not None:
                    runway_heading_values.add(runway_value % 360)
                if touchdown_data_index is None and runway_value is not None:
                    last_runway_heading = runway_value

                wind_speed_value = _finite_number(_cell(row, config.wind_speed_col))
                wind_direction_value = _finite_number(_cell(row, config.wind_direction_col))
                wind_pair_valid = (
                    wind_speed_value is not None
                    and wind_speed_value >= 0
                    and wind_direction_value is not None
                )
                if wind_pair_valid:
                    pair = (data_index, csv_row, float(wind_speed_value), float(wind_direction_value))
                    if touchdown_data_index is None:
                        last_wind_at_or_before = pair
                    elif data_index > touchdown_data_index and first_wind_after is None:
                        first_wind_after = pair

                radio_height = _finite_number(_cell(row, config.altitude_col))
                if not seen_above_2000 and radio_height is not None and radio_height > 2000:
                    seen_above_2000 = True

                if seen_above_2000 and touchdown_data_index is None:
                    aircraft_heading = _finite_number(_cell(row, config.aircraft_heading_col))
                    approach_rows.append(
                        (
                            data_index,
                            csv_row,
                            radio_height,
                            _numbers(row, *config.stick_cols),
                            _numbers(row, *config.rudder_cols),
                            aircraft_heading,
                            runway_value,
                        )
                    )
                    if radio_height is not None and radio_height < 0:
                        touchdown_data_index = data_index
                        touchdown_csv_row = csv_row
                        touchdown_aircraft_heading = aircraft_heading
                        touchdown_runway_heading = runway_value if runway_value is not None else last_runway_heading
                        if (
                            touchdown_aircraft_heading is not None
                            and touchdown_runway_heading is not None
                        ):
                            touchdown_angle = abs(
                                (
                                    (
                                        touchdown_aircraft_heading
                                        - touchdown_runway_heading
                                        + 180
                                    )
                                    % 360
                                )
                                - 180
                            )
                            landing_needs_wind = 6 < touchdown_angle <= 20

                current_pitch_rates = _numbers(row, *config.pitch_rate_cols)
                if takeoff_data_index is None:
                    group_1 = _numbers(row, *config.takeoff_group_1)
                    group_2 = _numbers(row, *config.takeoff_group_2)
                    count_1 = sum(value is not None and value > 2 for value in group_1)
                    count_2 = sum(value is not None and value > 2 for value in group_2)
                    if count_1 >= config.takeoff_min_1 and count_2 >= config.takeoff_min_2:
                        takeoff_data_index = data_index
                        takeoff_csv_row = csv_row
                        if previous_pitch_rates is not None:
                            pitch_window.extend(previous_pitch_rates)
                        pitch_window.extend(current_pitch_rates)
                        pitch_rows_remaining = 2
                        takeoff_baro = _finite_number(_cell(row, config.baro_altitude_col))
                    else:
                        previous_pitch_rates = current_pitch_rates
                elif pitch_rows_remaining > 0:
                    pitch_window.extend(current_pitch_rates)
                    pitch_rows_remaining -= 1

                wind_search_complete = (
                    touchdown_data_index is not None
                    and (
                        not landing_needs_wind
                        or (
                            last_wind_at_or_before is not None
                            and last_wind_at_or_before[0] == touchdown_data_index
                        )
                        or first_wind_after is not None
                    )
                )
                if (
                    touchdown_data_index is not None
                    and takeoff_data_index is not None
                    and pitch_rows_remaining == 0
                    and wind_search_complete
                    and len(runway_heading_values) >= 2
                ):
                    break
    except AnalysisCancelled:
        raise
    except Exception as exc:
        message = f"读取或解析CSV失败：{exc}"
        return FlightResult(
            person=person,
            path=path,
            aircraft_type=aircraft_type,
            flight_time=flight_time,
            flight_number=flight_number,
            metrics=_failed_metrics("CSV_READ_ERROR", message),
            crosswind_metric=MetricResult(
                True,
                value=None,
                error_code="NOT_APPLICABLE",
                error_message="未确认接地交叉角超限",
                applicable=False,
            ),
            warnings=warnings,
        )

    if short_rows:
        warnings.append(f"发现{short_rows}行字段数少于表头，缺失字段按无效值处理")
    if long_rows:
        warnings.append(f"发现{long_rows}行字段数多于表头，多余字段已忽略")

    metrics: dict[str, MetricResult] = {}
    if not seen_above_2000:
        landing_error = ("NO_ABOVE_2000", "D列未找到大于2000的有效高度")
    elif touchdown_data_index is None:
        landing_error = ("NO_TOUCHDOWN", "从首次大于2000后未找到D列首次负值")
    else:
        landing_error = None

    angle_difference: float | None = None
    signed_angle: float | None = None
    fifty_ft_csv_row: int | None = None
    rudder_region_samples: list[tuple[int, float]] = []
    reverse_rudder_value: float | None = None
    reverse_rudder_row: int | None = None
    if landing_error:
        metrics["stick"] = MetricResult(False, error_code=landing_error[0], error_message=landing_error[1])
        metrics["rudder"] = MetricResult(
            False, error_code=landing_error[0], error_message=landing_error[1]
        )
        metrics["cross_angle"] = MetricResult(False, error_code=landing_error[0], error_message=landing_error[1])
        metrics["reverse_rudder"] = MetricResult(
            False, error_code=landing_error[0], error_message=landing_error[1]
        )
    else:
        nearest_index = min(
            (index for index, item in enumerate(approach_rows) if item[2] is not None),
            key=lambda index: abs(float(approach_rows[index][2]) - 300),
        )
        stick_values = [
            value
            for _, _, _, row_values, _, _, _ in approach_rows[nearest_index:]
            for value in row_values
            if value is not None
        ]
        if stick_values:
            exceed = sum(abs(value) > 15 for value in stick_values)
            metrics["stick"] = MetricResult(True, {"beff": exceed, "ball": len(stick_values)})
        else:
            metrics["stick"] = MetricResult(
                False,
                error_code="NO_STICK_VALUES",
                error_message="300ft至接地区域没有有效杆量采样点",
            )

        if not rudder_columns_available:
            metrics["rudder"] = MetricResult(
                False,
                error_code="MISSING_RUDDER_COLUMNS",
                error_message=(
                    f"CSV列数不足，{aircraft_type.value}机型舵量需要"
                    f"{config.rudder_cols[0] + 1}至{config.rudder_cols[1] + 1}列"
                ),
            )
        else:
            valid_heights = [
                (index, float(item[2]))
                for index, item in enumerate(approach_rows)
                if item[2] is not None
            ]
            crossings = [
                (upper, lower)
                for upper, lower in zip(valid_heights, valid_heights[1:])
                if upper[1] >= 50 and lower[1] <= 50
            ]
            if not crossings:
                metrics["rudder"] = MetricResult(
                    False,
                    error_code="NO_50FT_CROSSING",
                    error_message="接地前未找到最终下降阶段由不低于50ft至不高于50ft的穿越",
                )
            else:
                upper, lower = crossings[-1]
                fifty_index = (
                    upper[0]
                    if abs(upper[1] - 50) <= abs(lower[1] - 50)
                    else lower[0]
                )
                fifty_ft_csv_row = approach_rows[fifty_index][1]
                rudder_region_samples = [
                    (csv_row, value)
                    for _, csv_row, _, _, row_values, _, _ in approach_rows[fifty_index:]
                    for value in row_values
                    if value is not None
                ]
                if not rudder_region_samples:
                    metrics["rudder"] = MetricResult(
                        False,
                        error_code="NO_RUDDER_VALUES",
                        error_message="50ft行至接地行区域没有有效舵量采样点",
                    )
                else:
                    rudder_exceed = sum(
                        abs(value) > 15 for _, value in rudder_region_samples
                    )
                    metrics["rudder"] = MetricResult(
                        True,
                        {
                            "reff": rudder_exceed,
                            "rall": len(rudder_region_samples),
                            "fifty_ft_row": fifty_ft_csv_row,
                            "touchdown_row": touchdown_csv_row,
                        },
                    )

        runway_heading_is_zero = (
            touchdown_runway_heading is not None
            and abs(touchdown_runway_heading % 360) < 1e-12
        )
        runway_heading_is_constant = (
            touchdown_runway_heading is not None and len(runway_heading_values) == 1
        )

        if runway_heading_is_zero:
            metrics["cross_angle"] = MetricResult(
                False,
                error_code="RUNWAY_HEADING_ZERO",
                error_message="跑道磁航向为0，判定为数据错误，不进行接地交叉角判断",
            )
        elif runway_heading_is_constant:
            metrics["cross_angle"] = MetricResult(
                False,
                error_code="RUNWAY_HEADING_CONSTANT",
                error_message="跑道磁航向列的有效数值始终保持不变，判定为数据错误，不进行接地交叉角判断",
            )
        elif touchdown_aircraft_heading is None:
            metrics["cross_angle"] = MetricResult(
                False, error_code="NO_AIRCRAFT_HEADING", error_message="接地行飞机磁航向无有效值"
            )
        elif touchdown_runway_heading is None:
            metrics["cross_angle"] = MetricResult(
                False, error_code="NO_RUNWAY_HEADING", error_message="接地行及其上方没有有效跑道磁航向"
            )
        else:
            signed_angle = (
                (touchdown_aircraft_heading - touchdown_runway_heading + 180) % 360
            ) - 180
            angle_difference = abs(signed_angle)
            if angle_difference > 20:
                metrics["cross_angle"] = MetricResult(
                    True,
                    value=None,
                    error_code="CROSS_ANGLE_OUT_OF_RANGE",
                    error_message=(
                        f"接地交叉角{angle_difference:.2f}°严格大于20°，"
                        "仅记录日志并排除接地交叉角判断"
                    ),
                    applicable=False,
                )
            else:
                metrics["cross_angle"] = MetricResult(True, angle_difference > 6)

        if runway_heading_is_zero:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="RUNWAY_HEADING_ZERO",
                error_message="跑道磁航向为0，判定为数据错误，不进行反向蹬舵判断",
            )
        elif runway_heading_is_constant:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="RUNWAY_HEADING_CONSTANT",
                error_message="跑道磁航向列的有效数值始终保持不变，判定为数据错误，不进行反向蹬舵判断",
            )
        elif angle_difference is not None and angle_difference > 20:
            metrics["reverse_rudder"] = MetricResult(
                True,
                value=None,
                error_code="CROSS_ANGLE_OUT_OF_RANGE",
                error_message=(
                    f"接地交叉角{angle_difference:.2f}°严格大于20°，"
                    "仅记录日志，不进行反向蹬舵判断"
                ),
                applicable=False,
            )
        elif not rudder_columns_available:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="MISSING_RUDDER_COLUMNS",
                error_message=(
                    f"CSV列数不足，{aircraft_type.value}机型舵量需要"
                    f"{config.rudder_cols[0] + 1}至{config.rudder_cols[1] + 1}列"
                ),
            )
        elif fifty_ft_csv_row is None:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="NO_50FT_CROSSING",
                error_message="接地前未找到最终下降阶段由不低于50ft至不高于50ft的穿越",
            )
        elif not rudder_region_samples:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="NO_RUDDER_VALUES",
                error_message="50ft行至接地行区域没有有效舵量采样点",
            )
        elif touchdown_aircraft_heading is None:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="NO_AIRCRAFT_HEADING",
                error_message="接地行飞机磁航向无有效值",
            )
        elif touchdown_runway_heading is None:
            metrics["reverse_rudder"] = MetricResult(
                False,
                error_code="NO_RUNWAY_HEADING",
                error_message="接地行及其上方没有有效跑道磁航向",
            )
        else:
            assert signed_angle is not None
            matching_samples: list[tuple[int, float]] = []
            if abs(signed_angle) > 3:
                matching_samples = [
                    (csv_row, value)
                    for csv_row, value in rudder_region_samples
                    if abs(value) > 1 and signed_angle * value < 0
                ]
            if matching_samples:
                reverse_rudder_row, reverse_rudder_value = matching_samples[0]
                for csv_row, value in matching_samples[1:]:
                    if abs(value) > abs(reverse_rudder_value):
                        reverse_rudder_row, reverse_rudder_value = csv_row, value
            metrics["reverse_rudder"] = MetricResult(
                True, reverse_rudder_value is not None
            )

    max_pitch_rate: float | None = None
    if takeoff_data_index is None:
        metrics["pitch_rate"] = MetricResult(
            False, error_code="NO_TAKEOFF_ROW", error_message="未找到满足条件的首次离地行"
        )
    else:
        valid_pitch_rates = [value for value in pitch_window if value is not None]
        if not valid_pitch_rates:
            metrics["pitch_rate"] = MetricResult(
                False, error_code="NO_PITCH_RATE_VALUES", error_message="离地行前1至后2行没有有效抬头速率"
            )
        else:
            max_pitch_rate = max(valid_pitch_rates)
            metrics["pitch_rate"] = MetricResult(True, max_pitch_rate > 3.5)

    pitch_metric = metrics["pitch_rate"]
    if pitch_metric.success and pitch_metric.value is True:
        if takeoff_baro is None:
            metrics["high_altitude"] = MetricResult(
                False, error_code="NO_BARO_ALTITUDE", error_message="离地行气压高度无有效值"
            )
        else:
            metrics["high_altitude"] = MetricResult(True, takeoff_baro > 8000)
    else:
        metrics["high_altitude"] = MetricResult(
            True,
            value=None,
            error_code="NOT_APPLICABLE",
            error_message="抬头速率未超限或抬头速率指标失败",
            applicable=False,
        )

    selected_wind: tuple[int, int, float, float] | None = None
    wind_speed: float | None = None
    wind_direction: float | None = None
    crosswind: float | None = None
    wind_source_row: int | None = None
    wind_row_offset: int | None = None
    cross_metric = metrics["cross_angle"]
    if cross_metric.success and cross_metric.value is True:
        candidates = [
            candidate
            for candidate in (last_wind_at_or_before, first_wind_after)
            if candidate is not None
        ]
        if touchdown_data_index is not None and candidates:
            selected_wind = min(
                candidates,
                key=lambda item: (
                    abs(item[0] - touchdown_data_index),
                    0 if item[0] <= touchdown_data_index else 1,
                ),
            )
        if selected_wind is None:
            crosswind_metric = MetricResult(
                False,
                error_code="NO_WIND_PAIR",
                error_message="CSV全部数据行均未找到同一行有效的风速和风向",
            )
        elif touchdown_runway_heading is None:
            crosswind_metric = MetricResult(
                False,
                error_code="NO_RUNWAY_HEADING",
                error_message="无法取得有效着陆跑道磁航向",
            )
        else:
            wind_data_index, wind_source_row, wind_speed, raw_wind_direction = selected_wind
            wind_direction = raw_wind_direction % 360
            wind_row_offset = wind_data_index - touchdown_data_index
            crosswind = calculate_crosswind(
                wind_speed,
                wind_direction,
                touchdown_runway_heading,
            )
            crosswind_metric = MetricResult(True, crosswind)
    else:
        crosswind_metric = MetricResult(
            True,
            value=None,
            error_code="NOT_APPLICABLE",
            error_message="接地交叉角未超限或交叉角指标失败",
            applicable=False,
        )

    return FlightResult(
        person=person,
        path=path,
        aircraft_type=aircraft_type,
        flight_time=flight_time,
        flight_number=flight_number,
        metrics=metrics,
        crosswind_metric=crosswind_metric,
        touchdown_row=touchdown_csv_row,
        fifty_ft_row=fifty_ft_csv_row,
        takeoff_row=takeoff_csv_row,
        aircraft_heading=touchdown_aircraft_heading,
        angle_difference=angle_difference,
        signed_angle=signed_angle,
        reverse_rudder_value=reverse_rudder_value,
        reverse_rudder_row=reverse_rudder_row,
        max_pitch_rate=max_pitch_rate,
        baro_altitude=takeoff_baro,
        runway_heading=touchdown_runway_heading,
        wind_speed=wind_speed,
        wind_direction=wind_direction,
        crosswind=crosswind,
        wind_source_row=wind_source_row,
        wind_row_offset=wind_row_offset,
        warnings=warnings,
    )


def _summarize_person(person: str, aircraft_category: str, flights: Iterable[FlightResult]) -> PersonSummary:
    flight_list = list(flights)
    stats = {key: MetricStats() for key in METRIC_LABELS}
    stick_exceed = 0
    stick_total = 0
    rudder_exceed = 0
    rudder_total = 0

    for flight in flight_list:
        for key in ("stick", "rudder", "cross_angle", "pitch_rate", "reverse_rudder"):
            metric = flight.metrics[key]
            item = stats[key]
            item.candidate += 1
            if not metric.applicable:
                continue
            if metric.success:
                item.valid += 1
                if key in ("stick", "rudder"):
                    exceed_key = "beff" if key == "stick" else "reff"
                    if metric.value[exceed_key] > 0:
                        item.triggered += 1
                elif metric.value is True:
                    item.triggered += 1
            else:
                item.failed += 1

        high_metric = flight.metrics["high_altitude"]
        if high_metric.applicable:
            stats["high_altitude"].candidate += 1
            if high_metric.success:
                stats["high_altitude"].valid += 1
                if high_metric.value is True:
                    stats["high_altitude"].triggered += 1
            else:
                stats["high_altitude"].failed += 1

        stick_metric = flight.metrics["stick"]
        if stick_metric.success:
            stick_exceed += int(stick_metric.value["beff"])
            stick_total += int(stick_metric.value["ball"])

        rudder_metric = flight.metrics["rudder"]
        if rudder_metric.success:
            rudder_exceed += int(rudder_metric.value["reff"])
            rudder_total += int(rudder_metric.value["rall"])

    cross = stats["cross_angle"]
    pitch = stats["pitch_rate"]
    high = stats["high_altitude"]
    return PersonSummary(
        person=person,
        aircraft_category=aircraft_category,
        stick_exceed_count=stick_exceed,
        stick_exceed_ratio=(stick_exceed / stick_total) if stick_total else 0.0,
        rudder_exceed_count=rudder_exceed,
        rudder_valid_sample_count=rudder_total,
        rudder_exceed_ratio=(rudder_exceed / rudder_total) if rudder_total else 0.0,
        cross_angle_count=cross.triggered,
        cross_angle_ratio=(cross.triggered / cross.valid) if cross.valid else 0.0,
        pitch_rate_count=pitch.triggered,
        pitch_rate_ratio=(pitch.triggered / pitch.valid) if pitch.valid else 0.0,
        high_altitude_count=high.triggered,
        high_altitude_ratio=(high.triggered / high.valid) if high.valid else 0.0,
        reverse_rudder_count=stats["reverse_rudder"].triggered,
        stats=stats,
    )


def _metric_log_record(flight: FlightResult, key: str, metric: MetricResult) -> LogRecord:
    if not metric.applicable:
        status = "NOT_APPLICABLE"
    else:
        status = "SUCCESS" if metric.success else "FAILED"
    value = ""
    if metric.value is not None:
        value = str(metric.value)
    if metric.success and key == "cross_angle" and flight.angle_difference is not None:
        value = f"触发={bool(metric.value)}；实际交叉角={flight.angle_difference:.2f}°"
    elif metric.success and key == "rudder":
        value = (
            f"Reff={metric.value['reff']}；Rall={metric.value['rall']}；"
            f"50ft行={metric.value['fifty_ft_row']}；"
            f"接地行={metric.value['touchdown_row']}"
        )
    elif metric.success and key == "pitch_rate" and flight.max_pitch_rate is not None:
        value = f"触发={bool(metric.value)}；最大抬头速率={flight.max_pitch_rate:.2f}°/s"
    elif metric.success and key == "high_altitude" and flight.baro_altitude is not None:
        value = f"高高原={bool(metric.value)}；气压高度={flight.baro_altitude:.2f} ft"
    elif metric.success and key == "reverse_rudder":
        direction = (
            "左交叉" if flight.signed_angle is not None and flight.signed_angle < 0 else "右交叉"
        )
        representative = (
            "" if flight.reverse_rudder_value is None else f"{flight.reverse_rudder_value:.2f}"
        )
        value = (
            f"触发={bool(metric.value)}；带方向交叉角={flight.signed_angle:.2f}°；"
            f"交叉方向={direction}；代表舵量={representative}；"
            f"代表舵量CSV行={flight.reverse_rudder_row or ''}；"
            f"50ft行={flight.fifty_ft_row or ''}；接地行={flight.touchdown_row or ''}"
        )
    message = metric.error_message
    if value:
        message = f"值={value}" + (f"；{message}" if message else "")
    direction = ""
    if key == "reverse_rudder" and flight.signed_angle is not None:
        if flight.signed_angle < 0:
            direction = "左交叉"
        elif flight.signed_angle > 0:
            direction = "右交叉"
    return LogRecord(
        record_type="DETAIL",
        timestamp=datetime.now(),
        person=flight.person,
        file_path=str(flight.path),
        aircraft_type=flight.aircraft_type.value,
        metric=METRIC_LABELS[key],
        status=status,
        error_code=metric.error_code,
        message=message,
        signed_angle=(
            f"{flight.signed_angle:.2f}"
            if key == "reverse_rudder" and flight.signed_angle is not None
            else ""
        ),
        cross_direction=direction,
        reverse_rudder_value=(
            f"{flight.reverse_rudder_value:.2f}"
            if key == "reverse_rudder" and flight.reverse_rudder_value is not None
            else ""
        ),
        reverse_rudder_row=(
            str(flight.reverse_rudder_row or "") if key == "reverse_rudder" else ""
        ),
        fifty_ft_row=(str(flight.fifty_ft_row or "") if key == "reverse_rudder" else ""),
        touchdown_row=(str(flight.touchdown_row or "") if key == "reverse_rudder" else ""),
        reverse_rudder_triggered=(
            str(bool(metric.value)) if key == "reverse_rudder" and metric.success else ""
        ),
        registration_raw=flight.registration_raw,
        registration=flight.registration,
        database_aircraft_type=flight.database_aircraft_type,
        aircraft_category=flight.aircraft_category,
        filename_engine_marker=CONFIGS[flight.aircraft_type].marker,
    )


def _crosswind_log_record(flight: FlightResult) -> LogRecord:
    metric = flight.crosswind_metric
    if not metric.applicable:
        status = "NOT_APPLICABLE"
    else:
        status = "SUCCESS" if metric.success else "FAILED"
    message = metric.error_message
    if metric.success and metric.value is not None:
        message = f"侧风分量={float(metric.value):.2f} kt"
    return LogRecord(
        record_type="DETAIL",
        timestamp=datetime.now(),
        person=flight.person,
        file_path=str(flight.path),
        aircraft_type=flight.aircraft_type.value,
        metric="侧风分量",
        status=status,
        error_code=metric.error_code,
        message=message,
        wind_source_row=str(flight.wind_source_row or ""),
        wind_row_offset=(
            str(flight.wind_row_offset)
            if flight.wind_row_offset is not None
            else ""
        ),
        crosswind=(
            f"{flight.crosswind:.2f}"
            if flight.crosswind is not None
            else ""
        ),
        registration_raw=flight.registration_raw,
        registration=flight.registration,
        database_aircraft_type=flight.database_aircraft_type,
        aircraft_category=flight.aircraft_category,
        filename_engine_marker=CONFIGS[flight.aircraft_type].marker,
    )


def analyze_selected_folder(
    path: str | Path,
    aircraft_database_path: str | Path,
    cancel_event: threading.Event | None = None,
    progress_callback: Callable[[ProgressEvent], None] | None = None,
) -> AnalysisResult:
    selected_folder = Path(path).resolve()
    database_path = Path(aircraft_database_path).resolve()
    cancel_event = cancel_event or threading.Event()
    callback = progress_callback or (lambda event: None)
    started_at = datetime.now()
    log_records = [
        LogRecord(
            record_type="RUN",
            timestamp=started_at,
            status="STARTED",
            message=f"所选数据目录={selected_folder}；飞机数据库={database_path}",
            aircraft_database=str(database_path),
        )
    ]
    people = discover_people(selected_folder)
    total_files = sum(len(items) for _, items in people)
    completed_files = 0
    all_flights: list[FlightResult] = []
    summaries: list[PersonSummary] = []
    cancelled = False

    callback(ProgressEvent("database", total=total_files, message="正在读取飞机数据库"))
    database = load_aircraft_database(database_path)
    if not database.success:
        ended_at = datetime.now()
        log_records.append(LogRecord(record_type="PREFLIGHT", timestamp=ended_at, status="FAILED", error_code=database.error_code, message=f"{database.error_message}；未执行指标分析，未生成Excel和DOCX", aircraft_database=str(database_path)))
        callback(ProgressEvent("preflight_failed", total=total_files, message=f"飞机数据库预检失败：{database.error_message}"))
        return AnalysisResult(selected_folder, database_path, started_at, ended_at, [], [], log_records, preflight_failed=True)

    callback(ProgressEvent("preflight", total=total_files, message="正在进行注册号预检"))
    preflight = preflight_candidates(people, database)
    for item in preflight.files:
        if item.status != "VALID":
            log_records.append(LogRecord(
                record_type="PREFLIGHT", timestamp=datetime.now(), person=item.person,
                file_path=str(item.path), aircraft_type=item.aircraft_type.value,
                status=item.status, error_code=item.error_code, message=item.message,
                aircraft_database=str(database_path), registration_raw=item.registration_raw,
                registration=item.registration, database_aircraft_type=item.database_aircraft_type,
                aircraft_category=item.aircraft_category,
                filename_engine_marker=CONFIGS[item.aircraft_type].marker,
                skipped="1" if item.status == "SKIPPED" else "0",
            ))
    if not preflight.success:
        ended_at = datetime.now()
        log_records.append(LogRecord(record_type="PREFLIGHT", timestamp=ended_at, status="FAILED", error_code="PREFLIGHT_BLOCKED", message="注册号预检失败；未执行指标分析，未生成Excel和DOCX", aircraft_database=str(database_path)))
        callback(ProgressEvent("preflight_failed", total=total_files, message=f"注册号预检失败：共{len(preflight.blocking_files)}个文件无法匹配"))
        return AnalysisResult(selected_folder, database_path, started_at, ended_at, [], [], log_records, preflight_failed=True)

    valid_by_person: dict[str, list[PreflightFileResult]] = {}
    for item in preflight.files:
        if item.status == "VALID":
            valid_by_person.setdefault(item.person, []).append(item)
    valid_count = sum(len(items) for items in valid_by_person.values())
    callback(ProgressEvent("started", total=valid_count, message=f"预检通过，发现{len(valid_by_person)}名人员、{valid_count}个有效数据文件"))
    for person_dir, _ in people:
        files = valid_by_person.get(person_dir.name, [])
        if not files:
            continue
        if cancel_event.is_set():
            cancelled = True
            break
        callback(
            ProgressEvent(
                "person", person=person_dir.name, completed=completed_files, total=valid_count,
                message=f"开始分析人员：{person_dir.name}",
            )
        )
        category_flights: dict[str, list[FlightResult]] = {}
        for item in files:
            file_path, aircraft_type = item.path, item.aircraft_type
            if cancel_event.is_set():
                cancelled = True
                break
            callback(
                ProgressEvent(
                    "file", person=person_dir.name, file_name=file_path.name,
                    completed=completed_files, total=valid_count, message=f"正在分析：{file_path.name}",
                )
            )
            try:
                flight = analyze_flight(person_dir.name, file_path, aircraft_type, cancel_event)
            except AnalysisCancelled:
                cancelled = True
                break
            flight.registration_raw = item.registration_raw
            flight.registration = item.registration
            flight.database_aircraft_type = item.database_aircraft_type
            flight.aircraft_category = item.aircraft_category
            category_flights.setdefault(item.aircraft_category, []).append(flight)
            all_flights.append(flight)
            for key, metric in flight.metrics.items():
                log_records.append(_metric_log_record(flight, key, metric))
            log_records.append(_crosswind_log_record(flight))
            if flight.flight_time is None:
                log_records.append(
                    LogRecord(
                        record_type="WARNING",
                        timestamp=datetime.now(),
                        person=flight.person,
                        file_path=str(flight.path),
                        aircraft_type=flight.aircraft_type.value,
                        metric="航段日期时间",
                        status="WARNING",
                        error_code="FLIGHT_TIME_PARSE_ERROR",
                        message="无法从文件名解析合法的14位航段日期时间，报告显示“无法解析”",
                    )
                )
            if flight.flight_number is None:
                log_records.append(
                    LogRecord(
                        record_type="WARNING",
                        timestamp=datetime.now(),
                        person=flight.person,
                        file_path=str(flight.path),
                        aircraft_type=flight.aircraft_type.value,
                        metric="航班号",
                        status="WARNING",
                        error_code="FLIGHT_NUMBER_PARSE_ERROR",
                        message="无法从文件名解析CA后跟数字的航班号，报告显示“无法解析”",
                    )
                )
            for warning in flight.warnings:
                log_records.append(
                    LogRecord(
                        record_type="WARNING", timestamp=datetime.now(), person=flight.person,
                        file_path=str(flight.path), aircraft_type=flight.aircraft_type.value,
                        status="WARNING", error_code="ROW_WIDTH_MISMATCH", message=warning,
                    )
                )
            completed_files += 1
            callback(
                ProgressEvent(
                    "progress", person=person_dir.name, file_name=file_path.name,
                    completed=completed_files, total=valid_count, message=f"已完成：{file_path.name}",
                )
            )
        for category in (AircraftCategory.A319.value, AircraftCategory.A320_A321.value, AircraftCategory.A321NEO.value):
            if category in category_flights:
                summaries.append(_summarize_person(person_dir.name, category, category_flights[category]))
        if cancelled:
            break

    for summary in summaries:
        skipped_for_group = sum(
            1 for item in preflight.files
            if item.status == "SKIPPED"
            and item.person == summary.person
            and item.aircraft_category == summary.aircraft_category
        )
        for stat in summary.stats.values():
            stat.skipped = skipped_for_group
        for key, item in summary.stats.items():
            summary_message = ""
            if key == "rudder":
                summary_message = (
                    f"Reff合计={summary.rudder_exceed_count}；"
                    f"Rall合计={summary.rudder_valid_sample_count}"
                )
                if summary.rudder_valid_sample_count == 0:
                    summary_message += "；该人员没有可用于舵量指标的有效采样点"
            log_records.append(
                LogRecord(
                    record_type="SUMMARY", timestamp=datetime.now(), person=summary.person,
                    metric=METRIC_LABELS[key], status="SUMMARY", candidate=str(item.candidate),
                    valid=str(item.valid), triggered=str(item.triggered), failed=str(item.failed),
                    skipped=str(item.skipped), message=summary_message,
                    aircraft_database=str(database_path), aircraft_category=summary.aircraft_category,
                )
            )

    summarized_groups = {(summary.person, summary.aircraft_category) for summary in summaries}
    skipped_groups: dict[tuple[str, str], int] = {}
    for item in preflight.files:
        if item.status == "SKIPPED" and item.aircraft_category:
            key = (item.person, item.aircraft_category)
            skipped_groups[key] = skipped_groups.get(key, 0) + 1
    for (person, category), skipped_count in skipped_groups.items():
        if (person, category) in summarized_groups:
            continue
        for key, label in METRIC_LABELS.items():
            log_records.append(LogRecord(
                record_type="SUMMARY", timestamp=datetime.now(), person=person,
                aircraft_category=category, metric=label, status="SUMMARY",
                candidate="0", valid="0", triggered="0", failed="0",
                skipped=str(skipped_count), aircraft_database=str(database_path),
                message="本人员及机型分类仅有预检跳过文件，未生成输出行",
            ))

    ended_at = datetime.now()
    failed_count = sum(1 for record in log_records if record.status == "FAILED")
    if cancelled:
        status = "CANCELLED"
        message = "用户取消分析，未生成新的Excel结果文件或DOCX报告"
    elif failed_count:
        status = "COMPLETED_WITH_ERRORS"
        message = f"分析完成，共记录{failed_count}项失败"
    else:
        status = "COMPLETED"
        message = "分析完成，未发现错误"
    log_records.append(LogRecord(record_type="ANALYSIS", timestamp=ended_at, status=status, message=message))
    callback(
        ProgressEvent(
            "cancelled" if cancelled else "analyzed",
            completed=completed_files,
            total=valid_count,
            message=message,
        )
    )
    return AnalysisResult(
        selected_folder=selected_folder,
        aircraft_database=database_path,
        started_at=started_at,
        ended_at=ended_at,
        summaries=summaries,
        flights=all_flights,
        log_records=log_records,
        cancelled=cancelled,
    )
