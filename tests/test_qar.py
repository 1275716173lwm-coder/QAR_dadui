from __future__ import annotations

import csv
import tempfile
import threading
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook

from qar_analyzer import (
    AircraftType,
    CONFIGS,
    calculate_crosswind,
    analyze_flight,
    analyze_selected_folder,
    parse_flight_number,
    parse_flight_time,
)
from qar_output import (
    LOG_HEADERS,
    WORKBOOK_HEADERS,
    OutputError,
    create_log_path,
    write_docx_report,
    write_log,
    write_workbook,
)


def set_value(row: list[str], index: int, value: object) -> None:
    row[index] = str(value)


def synthetic_cfm(
    path: Path,
    aircraft_heading: float = 359,
    runway_heading: float = 1,
    pitch_rate: float = 3.5,
    baro_altitude: float = 8000,
    wind_rows: dict[int, tuple[float, float]] | None = None,
    aircraft_type: AircraftType = AircraftType.CFM,
    altitudes: tuple[object, ...] = (2100, 300, 50, -1, -1),
    rudder_rows: dict[int, list[object]] | None = None,
) -> None:
    config = CONFIGS[aircraft_type]
    width = max(
        config.altitude_col,
        config.stick_cols[1],
        config.rudder_cols[1],
        config.aircraft_heading_col,
        config.runway_heading_col,
        config.pitch_rate_cols[1],
        config.baro_altitude_col,
        config.wind_speed_col,
        config.wind_direction_col,
    ) + 1
    header = [f"column_{index}" for index in range(width)]
    rows: list[list[str]] = []
    for altitude in altitudes:
        row = [""] * width
        if altitude != "":
            set_value(row, config.altitude_col, altitude)
        rows.append(row)

    for index in range(
        config.takeoff_group_1[0] + 1,
        config.takeoff_group_1[0] + 1 + config.takeoff_min_1,
    ):
        set_value(rows[0], index, 3)
    for index in range(
        config.takeoff_group_2[0],
        config.takeoff_group_2[0] + config.takeoff_min_2,
    ):
        set_value(rows[0], index, 3)
    set_value(rows[0], config.baro_altitude_col, baro_altitude)
    for row in rows:
        for index in range(config.pitch_rate_cols[0], config.pitch_rate_cols[1] + 1):
            set_value(row, index, pitch_rate)

    set_value(rows[1], config.stick_cols[0], 15)
    set_value(rows[1], config.stick_cols[0] + 1, -15)
    set_value(rows[1], config.stick_cols[0] + 2, 16)
    touchdown_index = next(
        index
        for index, altitude in enumerate(altitudes)
        if altitude != "" and float(altitude) < 0
    )
    set_value(rows[touchdown_index], config.aircraft_heading_col, aircraft_heading)
    set_value(rows[touchdown_index], config.runway_heading_col, runway_heading)

    if rudder_rows is None:
        rudder_rows = {
            touchdown_index - 1: [15, -15, 16],
            touchdown_index: [-16],
        }
    for row_index, values in rudder_rows.items():
        for offset, value in enumerate(values):
            set_value(rows[row_index], config.rudder_cols[0] + offset, value)

    if wind_rows is None:
        wind_rows = {touchdown_index: (10, 90)}
    for row_index, (speed, direction) in wind_rows.items():
        set_value(rows[row_index], config.wind_speed_col, speed)
        set_value(rows[row_index], config.wind_direction_col, direction)

    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerow(["units"] + [""] * (width - 1))
        writer.writerows(rows)


class QarAnalyzerTests(unittest.TestCase):
    def test_strict_thresholds_and_circular_heading(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "B-TEST_20260101010101___S1_E1_YZDCFM_CA0001.csv"
            synthetic_cfm(path)
            result = analyze_flight("测试", path, AircraftType.CFM, threading.Event())
            self.assertEqual(result.metrics["stick"].value["beff"], 1)
            self.assertEqual(result.metrics["rudder"].value["reff"], 2)
            self.assertEqual(result.metrics["rudder"].value["rall"], 4)
            self.assertFalse(result.metrics["cross_angle"].value)
            self.assertFalse(result.metrics["pitch_rate"].value)
            self.assertFalse(result.metrics["high_altitude"].applicable)
            self.assertFalse(result.crosswind_metric.applicable)

    def test_cross_angle_strictly_greater_than_six(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            six = folder / "B-A_20260101010101___S1_E1_YZDCFM_CA0001.csv"
            seven = folder / "B-B_20260101010102___S1_E1_YZDCFM_CA0002.csv"
            synthetic_cfm(six, 6, 0)
            synthetic_cfm(seven, 7, 0)
            self.assertFalse(
                analyze_flight("P", six, AircraftType.CFM, threading.Event())
                .metrics["cross_angle"].value
            )
            self.assertTrue(
                analyze_flight("P", seven, AircraftType.CFM, threading.Event())
                .metrics["cross_angle"].value
            )

    def test_cross_angle_over_twenty_is_invalid_for_cross_and_reverse_rudder(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            person = Path(temp) / "人员"
            person.mkdir()
            exact_twenty = person / "B-A_20260101010101_YZDCFM_CA0001.csv"
            over_twenty = person / "B-B_20260101010102_YZDCFM_CA0002.csv"
            synthetic_cfm(exact_twenty, aircraft_heading=20, runway_heading=0)
            synthetic_cfm(over_twenty, aircraft_heading=21, runway_heading=0)

            exact_result = analyze_flight(
                "人员", exact_twenty, AircraftType.CFM, threading.Event()
            )
            self.assertTrue(exact_result.metrics["cross_angle"].success)
            self.assertTrue(exact_result.metrics["cross_angle"].value)
            self.assertTrue(exact_result.metrics["reverse_rudder"].success)
            self.assertTrue(exact_result.metrics["reverse_rudder"].value)
            self.assertTrue(exact_result.crosswind_metric.success)

            invalid_result = analyze_flight(
                "人员", over_twenty, AircraftType.CFM, threading.Event()
            )
            self.assertEqual(invalid_result.angle_difference, 21)
            self.assertFalse(invalid_result.metrics["cross_angle"].success)
            self.assertEqual(
                invalid_result.metrics["cross_angle"].error_code,
                "CROSS_ANGLE_OUT_OF_RANGE",
            )
            self.assertFalse(invalid_result.metrics["reverse_rudder"].success)
            self.assertEqual(
                invalid_result.metrics["reverse_rudder"].error_code,
                "CROSS_ANGLE_OUT_OF_RANGE",
            )
            self.assertTrue(invalid_result.metrics["rudder"].success)
            self.assertFalse(invalid_result.crosswind_metric.applicable)

            result = analyze_selected_folder(temp)
            summary = result.summaries[0]
            self.assertEqual(summary.cross_angle_count, 1)
            self.assertEqual(summary.cross_angle_ratio, 1.0)
            self.assertEqual(summary.reverse_rudder_count, 1)
            self.assertEqual(
                (
                    summary.stats["cross_angle"].candidate,
                    summary.stats["cross_angle"].valid,
                    summary.stats["cross_angle"].triggered,
                    summary.stats["cross_angle"].failed,
                ),
                (2, 1, 1, 1),
            )
            self.assertEqual(
                (
                    summary.stats["reverse_rudder"].candidate,
                    summary.stats["reverse_rudder"].valid,
                    summary.stats["reverse_rudder"].triggered,
                    summary.stats["reverse_rudder"].failed,
                ),
                (2, 1, 1, 1),
            )

    def test_crosswind_direction_examples(self) -> None:
        self.assertAlmostEqual(calculate_crosswind(10, 90, 0), 10)
        self.assertAlmostEqual(calculate_crosswind(10, 270, 0), -10)
        self.assertEqual(calculate_crosswind(10, 0, 0), 0)
        self.assertEqual(calculate_crosswind(10, 180, 0), 0)
        self.assertAlmostEqual(calculate_crosswind(10, -90, 0), -10)
        self.assertAlmostEqual(calculate_crosswind(10, 450, 0), 10)

    def test_nearest_wind_pair_tie_prefers_upward_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "B-TEST_20260101010101___S1_E1_YZDCFM_CA0001.csv"
            synthetic_cfm(
                path,
                aircraft_heading=10,
                runway_heading=0,
                wind_rows={2: (10, 90), 4: (20, 270)},
            )
            result = analyze_flight("测试", path, AircraftType.CFM, threading.Event())
            self.assertTrue(result.crosswind_metric.success)
            self.assertEqual(result.wind_source_row, 5)
            self.assertEqual(result.wind_row_offset, -1)
            self.assertAlmostEqual(result.crosswind or 0, 10)

    def test_rudder_columns_for_all_aircraft_types(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            markers = {
                AircraftType.CFM: "YZDCFM",
                AircraftType.LEAP: "YZDLEAP",
                AircraftType.PW: "YZDPW",
            }
            for aircraft_type, marker in markers.items():
                path = folder / f"B-X_20260101010101_{marker}_CA0001.csv"
                synthetic_cfm(path, aircraft_type=aircraft_type)
                result = analyze_flight("测试", path, aircraft_type, threading.Event())
                self.assertTrue(result.metrics["rudder"].success)
                self.assertEqual(result.metrics["rudder"].value["reff"], 2)
                self.assertEqual(result.metrics["rudder"].value["rall"], 4)

    def test_rudder_uses_last_50ft_crossing_and_skips_invalid_height(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "B-X_20260101010101_YZDCFM_CA0001.csv"
            synthetic_cfm(
                path,
                altitudes=(2100, 50, 500, 60, "", 40, -1, -1),
                rudder_rows={1: [20], 3: [15], 5: [-15], 6: [14]},
            )
            metric = analyze_flight(
                "测试", path, AircraftType.CFM, threading.Event()
            ).metrics["rudder"]
            self.assertTrue(metric.success)
            self.assertEqual(metric.value["fifty_ft_row"], 6)
            self.assertEqual(metric.value["reff"], 0)
            self.assertEqual(metric.value["rall"], 3)

    def test_rudder_failure_is_independent_when_columns_are_missing(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "B-X_20260101010101_YZDCFM_CA0001.csv"
            synthetic_cfm(path)
            with path.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.reader(handle))
            shortened = [row[: CONFIGS[AircraftType.CFM].rudder_cols[0]] for row in rows]
            with path.open("w", encoding="utf-8", newline="") as handle:
                csv.writer(handle).writerows(shortened)

            result = analyze_flight("测试", path, AircraftType.CFM, threading.Event())
            self.assertFalse(result.metrics["rudder"].success)
            self.assertEqual(
                result.metrics["rudder"].error_code,
                "MISSING_RUDDER_COLUMNS",
            )
            self.assertTrue(result.metrics["stick"].success)
            self.assertFalse(result.metrics["reverse_rudder"].success)

    def test_reverse_rudder_direction_threshold_and_representative_sample(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            left = folder / "B-L_20260101010101_YZDCFM_CA0001.csv"
            right = folder / "B-R_20260101010102_YZDCFM_CA0002.csv"
            exact_three = folder / "B-E_20260101010103_YZDCFM_CA0003.csv"
            wrong_direction = folder / "B-W_20260101010104_YZDCFM_CA0004.csv"
            zero_only = folder / "B-Z_20260101010105_YZDCFM_CA0005.csv"

            synthetic_cfm(
                left,
                aircraft_heading=357,
                runway_heading=1,
                rudder_rows={2: [0, 0.8], 3: [0.8]},
            )
            synthetic_cfm(
                right,
                aircraft_heading=5,
                runway_heading=1,
                rudder_rows={2: [-0.1]},
            )
            synthetic_cfm(
                exact_three,
                aircraft_heading=358,
                runway_heading=1,
                rudder_rows={2: [20]},
            )
            synthetic_cfm(
                wrong_direction,
                aircraft_heading=357,
                runway_heading=1,
                rudder_rows={2: [-20, 0]},
            )
            synthetic_cfm(
                zero_only,
                aircraft_heading=357,
                runway_heading=1,
                rudder_rows={2: [0]},
            )

            left_result = analyze_flight(
                "测试", left, AircraftType.CFM, threading.Event()
            )
            self.assertEqual(left_result.signed_angle, -4)
            self.assertTrue(left_result.metrics["reverse_rudder"].value)
            self.assertEqual(left_result.reverse_rudder_value, 0.8)
            self.assertEqual(left_result.reverse_rudder_row, 5)
            self.assertEqual(left_result.fifty_ft_row, 5)

            right_result = analyze_flight(
                "测试", right, AircraftType.CFM, threading.Event()
            )
            self.assertEqual(right_result.signed_angle, 4)
            self.assertTrue(right_result.metrics["reverse_rudder"].value)
            self.assertEqual(right_result.reverse_rudder_value, -0.1)

            exact_result = analyze_flight(
                "测试", exact_three, AircraftType.CFM, threading.Event()
            )
            self.assertEqual(exact_result.signed_angle, -3)
            self.assertFalse(exact_result.metrics["reverse_rudder"].value)

            wrong_result = analyze_flight(
                "测试", wrong_direction, AircraftType.CFM, threading.Event()
            )
            self.assertFalse(wrong_result.metrics["reverse_rudder"].value)

            zero_result = analyze_flight(
                "测试", zero_only, AircraftType.CFM, threading.Event()
            )
            self.assertFalse(zero_result.metrics["reverse_rudder"].value)

    def test_reverse_rudder_summary_counts_each_flight_once(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            person = Path(temp) / "人员"
            person.mkdir()
            synthetic_cfm(
                person / "B-L_20260101010101_YZDCFM_CA0001.csv",
                aircraft_heading=357,
                runway_heading=1,
                rudder_rows={2: [0.1, 0.2, 0.3], 3: [0.4]},
            )
            synthetic_cfm(
                person / "B-N_20260101010102_YZDCFM_CA0002.csv",
                aircraft_heading=358,
                runway_heading=1,
                rudder_rows={2: [20]},
            )

            result = analyze_selected_folder(temp)
            self.assertEqual(result.summaries[0].reverse_rudder_count, 1)
            stats = result.summaries[0].stats["reverse_rudder"]
            self.assertEqual((stats.candidate, stats.valid, stats.triggered, stats.failed), (2, 2, 1, 0))

    def test_no_valid_rudder_values_leave_ratio_empty(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            person = Path(temp) / "人员"
            person.mkdir()
            path = person / "B-X_20260101010101_YZDCFM_CA0001.csv"
            synthetic_cfm(path, rudder_rows={})

            result = analyze_selected_folder(temp)
            metric = result.flights[0].metrics["rudder"]
            self.assertFalse(metric.success)
            self.assertEqual(metric.error_code, "NO_RUDDER_VALUES")
            self.assertEqual(result.summaries[0].rudder_exceed_count, 0)
            self.assertEqual(result.summaries[0].rudder_valid_sample_count, 0)
            self.assertIsNone(result.summaries[0].rudder_exceed_ratio)
            self.assertEqual(result.summaries[0].stats["rudder"].failed, 1)
            self.assertTrue(
                any(
                    record.metric == "50ft以下舵量"
                    and "没有可用于舵量指标的有效采样点" in record.message
                    for record in result.log_records
                )
            )

    def test_pitch_event_records_peak_and_high_altitude(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "B-TEST_20260101010101___S1_E1_YZDCFM_CA0001.csv"
            synthetic_cfm(path, pitch_rate=3.6, baro_altitude=8001)
            result = analyze_flight("测试", path, AircraftType.CFM, threading.Event())
            self.assertTrue(result.metrics["pitch_rate"].value)
            self.assertAlmostEqual(result.max_pitch_rate or 0, 3.6)
            self.assertTrue(result.metrics["high_altitude"].value)

    def test_metadata_parsing_uses_first_legal_values(self) -> None:
        name = "B-X_20261301010101_20260102030405_YZDCFM_CA1234_CA9999.csv"
        self.assertEqual(parse_flight_time(name), datetime(2026, 1, 2, 3, 4, 5))
        self.assertEqual(parse_flight_number(name, AircraftType.CFM), "CA1234")

    def test_cancel_does_not_analyze_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            person = Path(temp) / "人员"
            person.mkdir()
            synthetic_cfm(person / "B-A_20260101010101___S1_E1_YZDCFM_CA0001.csv")
            event = threading.Event()
            event.set()
            result = analyze_selected_folder(temp, event)
            self.assertTrue(result.cancelled)
            self.assertEqual(result.flights, [])

    def test_timestamped_outputs_docx_and_log(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            person = root / "input" / "甲"
            person.mkdir(parents=True)
            synthetic_cfm(
                person / "B-A_20260101010101___S1_E1_YZDCFM_CA0001.csv",
                aircraft_heading=10,
                runway_heading=0,
                pitch_rate=3.6,
                baro_altitude=8001,
            )
            result = analyze_selected_folder(root / "input")

            template = root / "template.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for column in range(1, 13):
                sheet.cell(1, column, f"H{column}")
            workbook.save(template)
            workbook.close()

            output_dir = root / "result"
            fixed_save_time = datetime(2026, 1, 2, 3, 4, 5)
            with patch("qar_output._local_now", return_value=fixed_save_time):
                output = write_workbook(template, output_dir, result.summaries)
                second_output = write_workbook(template, output_dir, result.summaries)
            self.assertEqual(output.name, "qar_20260102_030405.xlsx")
            self.assertEqual(second_output.name, "qar_20260102_030405_1.xlsx")

            saved = load_workbook(output, data_only=False)
            sheet = saved.active
            self.assertEqual(
                [sheet.cell(1, column).value for column in range(1, 13)],
                WORKBOOK_HEADERS,
            )
            self.assertEqual(sheet.cell(2, 1).value, "甲")
            self.assertEqual(sheet.cell(2, 4).value, 2)
            self.assertEqual(sheet.cell(2, 5).value, 0.5)
            self.assertEqual(sheet.cell(2, 5).number_format, "0.00%")
            self.assertEqual(sheet.cell(2, 12).value, 1)
            self.assertEqual(sheet.cell(2, 3).number_format, "0.00%")
            self.assertTrue(sheet.cell(1, 1).alignment.wrap_text)
            self.assertGreaterEqual(sheet.row_dimensions[1].height or 0, 96)
            saved.close()

            report = write_docx_report(output_dir, result)
            second_report = write_docx_report(output_dir, result)
            self.assertTrue(report.name.startswith("qar_report_"))
            self.assertTrue(second_report.stem.endswith("_1"))
            document = Document(report)
            paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            table_text = "\n".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("甲", paragraph_text)
            self.assertIn("接地交叉角事件", paragraph_text)
            self.assertIn("抬头速率事件", paragraph_text)
            self.assertIn("疑似接地前反向蹬舵事件", paragraph_text)
            self.assertIn("右交叉10.00°", paragraph_text)
            self.assertIn("右舵-16.00", paragraph_text)
            self.assertIn("代表舵量数据所在CSV行", paragraph_text)
            self.assertNotIn("暂未分析", table_text)
            self.assertIn("50.00%", table_text)
            self.assertNotIn("侧风分量", paragraph_text.split("抬头速率事件", 1)[1])

            log_path = write_log(output_dir, result)
            content = log_path.read_text(encoding="utf-8")
            self.assertEqual(content.splitlines()[0].split("\t"), LOG_HEADERS)
            self.assertIn("Reff=2；Rall=4；50ft行=5；接地行=6", content)
            self.assertIn("Reff合计=2；Rall合计=4", content)
            self.assertIn("疑似接地前反向蹬舵", content)
            self.assertIn("带方向交叉角=10.00°", content)
            self.assertTrue(log_path.name.startswith("qar_log_"))
            second_log = create_log_path(output_dir, result.started_at)
            self.assertNotEqual(log_path, second_log)
            self.assertEqual(list(output_dir.glob(".*.tmp.xlsx")), [])
            self.assertEqual(list(output_dir.glob(".*.tmp.docx")), [])

    def test_cancelled_result_cannot_create_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            result = analyze_selected_folder(temp, threading.Event())
            result.cancelled = True
            with self.assertRaises(OutputError):
                write_docx_report(Path(temp) / "result", result)
            self.assertEqual(list((Path(temp) / "result").glob("*.docx")), [])


if __name__ == "__main__":
    unittest.main()
