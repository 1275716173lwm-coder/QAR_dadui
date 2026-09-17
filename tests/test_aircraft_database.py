from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from qar_analyzer import (
    AircraftCategory,
    AircraftType,
    SUPPORTED_AIRCRAFT,
    analyze_selected_folder,
    load_aircraft_database,
)
from qar_output import WORKBOOK_HEADERS, write_docx_report, write_workbook
from tests.test_qar import synthetic_cfm


def make_database(path: Path, rows: list[tuple[str, str]], headers: tuple[str, str] = ("Registration", "A/C TYPE")) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    workbook.close()
    return path


class AircraftDatabaseTests(unittest.TestCase):
    def test_all_supported_aircraft_mappings(self) -> None:
        expected = {
            "A319CFM": (AircraftType.CFM, AircraftCategory.A319),
            "A319LEAP": (AircraftType.LEAP, AircraftCategory.A319),
            "A320CFM": (AircraftType.CFM, AircraftCategory.A320_A321),
            "A321CFM": (AircraftType.CFM, AircraftCategory.A320_A321),
            "A320PW": (AircraftType.PW, AircraftCategory.A320_A321),
            "A320LEAP": (AircraftType.LEAP, AircraftCategory.A320_A321),
            "A321LEAP": (AircraftType.LEAP, AircraftCategory.A321NEO),
            "A321PW": (AircraftType.PW, AircraftCategory.A321NEO),
        }
        self.assertEqual(SUPPORTED_AIRCRAFT, expected)

    def test_database_normalizes_registration_and_detects_conflict(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            valid = make_database(root / "valid.xlsx", [(" b-a1z9 ", " a319cfm ")])
            loaded = load_aircraft_database(valid)
            self.assertTrue(loaded.success)
            self.assertEqual(loaded.registrations["B-A1Z9"], "A319CFM")

            conflict = make_database(root / "conflict.xlsx", [("B-A1Z9", "A319CFM"), (" b-a1z9 ", "A320CFM")])
            loaded = load_aircraft_database(conflict)
            self.assertFalse(loaded.success)
            self.assertIn("重复注册号", loaded.error_message)

    def test_missing_or_unmatched_registration_blocks_all_analysis(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            person = root / "input" / "甲"
            person.mkdir(parents=True)
            synthetic_cfm(person / "NO_REG_20260101010101_YZDCFM_CA1.csv")
            synthetic_cfm(person / "B-Z9Z9_20260101010102_YZDCFM_CA2.csv")
            database = make_database(root / "reg.xlsx", [("B-AAAA", "A319CFM")])
            result = analyze_selected_folder(root / "input", database)
            self.assertTrue(result.preflight_failed)
            self.assertEqual(result.flights, [])
            self.assertEqual(result.summaries, [])
            codes = {record.error_code for record in result.log_records}
            self.assertIn("REGISTRATION_MISSING", codes)
            self.assertIn("REGISTRATION_NOT_FOUND", codes)
            self.assertTrue(any("未执行指标分析，未生成Excel和DOCX" in record.message for record in result.log_records))

    def test_unsupported_and_engine_conflict_are_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            person = root / "input" / "甲"
            person.mkdir(parents=True)
            synthetic_cfm(person / "B-AAAA_20260101010101_YZDCFM_CA1.csv")
            synthetic_cfm(person / "B-BBBB_20260101010102_YZDCFM_CA2.csv")
            synthetic_cfm(person / "B-CCCC_20260101010103_YZDCFM_CA3.csv")
            database = make_database(root / "reg.xlsx", [
                ("B-AAAA", "A319CFM"),
                ("B-BBBB", "BOEING737"),
                ("B-CCCC", "A319LEAP"),
            ])
            result = analyze_selected_folder(root / "input", database)
            self.assertFalse(result.preflight_failed)
            self.assertEqual(len(result.flights), 1)
            self.assertEqual(result.flights[0].registration, "B-AAAA")
            skipped_codes = {record.error_code for record in result.log_records if record.status == "SKIPPED"}
            self.assertEqual(skipped_codes, {"AIRCRAFT_TYPE_UNSUPPORTED", "ENGINE_MARKER_CONFLICT"})

    def test_category_outputs_are_isolated_merged_and_three_level_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            person = root / "input" / "甲"
            person.mkdir(parents=True)
            synthetic_cfm(person / "B-AAAA_20260101010101_YZDCFM_CA1.csv")
            synthetic_cfm(person / "B-BBBB_20260101010102_YZDCFM_CA2.csv")
            synthetic_cfm(person / "B-CCCC_20260101010103_YZDCFM_CA3.csv")
            database = make_database(root / "reg.xlsx", [
                (" b-aaaa ", "A319CFM"),
                ("B-BBBB", "A320CFM"),
                ("B-CCCC", "A321CFM"),
            ])
            result = analyze_selected_folder(root / "input", database)
            self.assertEqual([item.aircraft_category for item in result.summaries], ["A319", "A320、A321"])
            self.assertEqual([item.stats["stick"].candidate for item in result.summaries], [1, 2])

            template = root / "template.xlsx"
            workbook = Workbook()
            for column in range(1, 14):
                workbook.active.cell(1, column, f"H{column}")
            workbook.save(template)
            workbook.close()
            output_dir = root / "result"
            output = write_workbook(template, output_dir, result.summaries)
            saved = load_workbook(output)
            sheet = saved.active
            self.assertEqual([sheet.cell(1, col).value for col in range(1, 14)], WORKBOOK_HEADERS)
            self.assertIn("A2:A3", {str(item) for item in sheet.merged_cells.ranges})
            self.assertEqual([sheet.cell(row, 2).value for row in (2, 3)], ["A319", "A320、A321"])
            saved.close()

            report = write_docx_report(output_dir, result)
            document = Document(report)
            headings = [(p.style.name, p.text) for p in document.paragraphs if p.style and p.style.name.startswith("Heading")]
            self.assertIn(("Heading 1", "甲"), headings)
            self.assertIn(("Heading 2", "A319"), headings)
            self.assertIn(("Heading 2", "A320、A321"), headings)
            self.assertIn(("Heading 3", "指标汇总"), headings)


if __name__ == "__main__":
    unittest.main()
