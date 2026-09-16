from __future__ import annotations

import os
import uuid
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook

from qar_analyzer import AnalysisResult, FlightResult, LogRecord, PersonSummary


class OutputError(RuntimeError):
    pass


LOG_HEADERS = [
    "记录类型", "时间", "人员", "文件完整路径", "机型", "指标", "状态", "错误代码",
    "说明", "候选文件数", "有效文件数", "触发文件数", "失败文件数",
    "Excel结果文件", "DOCX状态", "风数据来源行", "风数据行差", "侧风分量kt",
    "带方向接地交叉角", "交叉方向", "反向蹬舵代表值", "代表值CSV行",
    "50ft行", "接地行", "反向蹬舵触发",
]

WORKBOOK_HEADERS = [
    "姓名",
    "低空 300ft AGL 以下杆量超过 15 个单位采样点数",
    "低空 300ft AGL 以下杆量超过 15 个单位采样点占比",
    "低空 50ft AGL 以下舵量超过 15 个单位次数",
    "低空 50ft AGL 以下舵量超过 15 个单位占比",
    "接地交叉角大于 6° 次数",
    "接地交叉角大于 6° 占比",
    "PF 抬头速率大于 3.5°/s 次数",
    "PF 抬头速率大于 3.5°/s 占比",
    "抬头速率大中高高原次数",
    "抬头速率大中高高原占比",
    "疑似接地前反向蹬舵次数",
]


def _safe_text(value: object) -> str:
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _local_now() -> datetime:
    return datetime.now()


def _numbered_path(output_dir: Path, stem: str, suffix: str) -> Path:
    candidate = output_dir / f"{stem}{suffix}"
    counter = 1
    while candidate.exists():
        candidate = output_dir / f"{stem}_{counter}{suffix}"
        counter += 1
    return candidate


def _publish_without_overwrite(temporary_path: Path, target_path: Path, label: str) -> Path:
    candidate = target_path
    counter = 1
    while True:
        try:
            os.link(temporary_path, candidate)
            temporary_path.unlink()
            return candidate
        except FileExistsError:
            candidate = target_path.with_name(f"{target_path.stem}_{counter}{target_path.suffix}")
            counter += 1
        except PermissionError as exc:
            raise OutputError(f"无法发布{label}，请检查result文件夹权限：{exc}") from exc
        except OSError as exc:
            raise OutputError(f"无法发布{label}：{exc}") from exc


def create_log_path(output_dir: Path, started_at: datetime) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"qar_log_{started_at:%Y%m%d_%H%M%S}"
    return _numbered_path(output_dir, stem, ".txt")


def write_log(output_dir: Path, result: AnalysisResult, extra_records: list[LogRecord] | None = None) -> Path:
    path = create_log_path(output_dir, result.started_at)
    records = [*result.log_records, *(extra_records or [])]
    with path.open("x", encoding="utf-8", newline="") as handle:
        handle.write("\t".join(LOG_HEADERS) + "\n")
        for record in records:
            row = [
                record.record_type,
                record.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                record.person,
                record.file_path,
                record.aircraft_type,
                record.metric,
                record.status,
                record.error_code,
                record.message,
                record.candidate,
                record.valid,
                record.triggered,
                record.failed,
                record.excel_file,
                record.docx_status,
                record.wind_source_row,
                record.wind_row_offset,
                record.crosswind,
                record.signed_angle,
                record.cross_direction,
                record.reverse_rudder_value,
                record.reverse_rudder_row,
                record.fifty_ft_row,
                record.touchdown_row,
                record.reverse_rudder_triggered,
            ]
            handle.write("\t".join(_safe_text(value) for value in row) + "\n")
    return path


def _write_summary_row(sheet, row: int, summary: PersonSummary) -> None:
    values = [
        summary.person,
        summary.stick_exceed_count,
        summary.stick_exceed_ratio,
        summary.rudder_exceed_count,
        summary.rudder_exceed_ratio,
        summary.cross_angle_count,
        summary.cross_angle_ratio,
        summary.pitch_rate_count,
        summary.pitch_rate_ratio,
        summary.high_altitude_count,
        summary.high_altitude_ratio,
        summary.reverse_rudder_count,
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(row=row, column=column, value=value)
    for column in (3, 5, 7, 9, 11):
        sheet.cell(row=row, column=column).number_format = "0.00%"


def write_workbook(template_path: Path, output_dir: Path, summaries: list[PersonSummary]) -> Path:
    if not template_path.exists() or not template_path.is_file():
        raise OutputError(f"结果模板不存在：{template_path}")
    output_dir.mkdir(parents=True, exist_ok=True)
    temporary_path = output_dir / f".qar_{uuid.uuid4().hex}.tmp.xlsx"
    workbook = None

    try:
        workbook = load_workbook(template_path)
        if not workbook.worksheets:
            raise OutputError("结果模板不包含工作表")
        sheet = workbook.worksheets[0]
        if sheet.max_column < 12:
            raise OutputError(f"结果模板至少需要12列，当前只有{sheet.max_column}列")
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for column, header in enumerate(WORKBOOK_HEADERS, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            alignment.horizontal = "center"
            alignment.vertical = "center"
            cell.alignment = alignment
        sheet.row_dimensions[1].height = max(sheet.row_dimensions[1].height or 0, 96)
        for row, summary in enumerate(summaries, start=2):
            _write_summary_row(sheet, row, summary)
        workbook.save(temporary_path)
        workbook.close()
        workbook = None

        validation_workbook = load_workbook(temporary_path, read_only=True, data_only=False)
        try:
            if not validation_workbook.worksheets:
                raise OutputError("临时结果工作簿校验失败：没有工作表")
            validation_sheet = validation_workbook.worksheets[0]
            if validation_sheet.max_column < 12:
                raise OutputError("临时结果工作簿校验失败：列数不足12列")
        finally:
            validation_workbook.close()

        saved_at = _local_now()
        target_path = output_dir / f"qar_{saved_at:%Y%m%d_%H%M%S}.xlsx"
        return _publish_without_overwrite(temporary_path, target_path, "Excel结果文件")
    except OutputError:
        raise
    except Exception as exc:
        raise OutputError(f"生成结果表失败：{exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass


def _set_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), name)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    style_settings = {
        "Normal": (10.5, False, 6),
        "Title": (20, True, 12),
        "Heading 1": (15, True, 8),
        "Heading 2": (12, True, 6),
        "Heading 3": (11, True, 4),
    }
    for style_name, (size, bold, after) in style_settings.items():
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(after)
        paragraph_properties = style._element.get_or_add_pPr()
        paragraph_borders = paragraph_properties.find(qn("w:pBdr"))
        if paragraph_borders is not None:
            paragraph_properties.remove(paragraph_borders)
        if style_name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True


def _set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_edges_and_padding(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")

    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin_name in ("top", "left", "bottom", "right"):
        margin = margins.find(qn(f"w:{margin_name}"))
        if margin is None:
            margin = OxmlElement(f"w:{margin_name}")
            margins.append(margin)
        margin.set(qn("w:w"), "100")
        margin.set(qn("w:type"), "dxa")


def _add_two_column_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_cells = table.rows[0].cells
    header_cells[0].text = "指标"
    header_cells[1].text = "结果"
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_properties.append(OxmlElement("w:tblHeader"))

    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(4.65 if column_index == 0 else 2.2)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_edges_and_padding(cell)
            if row_index == 0:
                _set_cell_fill(cell, "1F4E78")
            elif row_index % 2 == 0:
                _set_cell_fill(cell, "F2F6FA")
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_font(
                        run,
                        "Microsoft YaHei",
                        9.5,
                        bold=(row_index == 0),
                    )
                    if row_index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph()


def _format_ratio(value: float | None) -> str:
    return "无有效数据" if value is None else f"{value:.2%}"


def _summary_rows(summary: PersonSummary) -> list[tuple[str, str]]:
    return [
        (WORKBOOK_HEADERS[0], summary.person),
        (WORKBOOK_HEADERS[1], str(summary.stick_exceed_count)),
        (WORKBOOK_HEADERS[2], _format_ratio(summary.stick_exceed_ratio)),
        (WORKBOOK_HEADERS[3], str(summary.rudder_exceed_count)),
        (WORKBOOK_HEADERS[4], _format_ratio(summary.rudder_exceed_ratio)),
        (WORKBOOK_HEADERS[5], str(summary.cross_angle_count)),
        (WORKBOOK_HEADERS[6], _format_ratio(summary.cross_angle_ratio)),
        (WORKBOOK_HEADERS[7], str(summary.pitch_rate_count)),
        (WORKBOOK_HEADERS[8], _format_ratio(summary.pitch_rate_ratio)),
        (WORKBOOK_HEADERS[9], str(summary.high_altitude_count)),
        (WORKBOOK_HEADERS[10], _format_ratio(summary.high_altitude_ratio)),
        (WORKBOOK_HEADERS[11], str(summary.reverse_rudder_count)),
    ]


def _add_labeled_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}：")
    _set_font(label_run, "Microsoft YaHei", 10.5, bold=True)
    value_run = paragraph.add_run(value)
    _set_font(value_run, "Microsoft YaHei", 10.5)


def _keep_event_block_together(document: Document, start_index: int) -> None:
    paragraphs = document.paragraphs[start_index:]
    for paragraph in paragraphs[:-1]:
        paragraph.paragraph_format.keep_with_next = True


def _display_time(flight: FlightResult) -> str:
    return (
        flight.flight_time.strftime("%Y-%m-%d %H:%M:%S")
        if flight.flight_time is not None
        else "无法解析"
    )


def _display_number(value: float | None, suffix: str = "", digits: int = 2) -> str:
    return "无法计算" if value is None else f"{value:.{digits}f}{suffix}"


def _high_altitude_status(flight: FlightResult) -> str:
    metric = flight.metrics["high_altitude"]
    if not metric.applicable or not metric.success:
        return "无法判断"
    return "是" if metric.value is True else "否"


def _add_cross_angle_events(document: Document, flights: list[FlightResult]) -> None:
    events = [
        flight
        for flight in flights
        if flight.metrics["cross_angle"].success
        and flight.metrics["cross_angle"].value is True
    ]
    if not events:
        document.add_paragraph("无")
        return

    for event_index, flight in enumerate(events, start=1):
        start_index = len(document.paragraphs)
        document.add_heading(f"事件 {event_index}", level=3)
        _add_labeled_paragraph(document, "航段日期时间", _display_time(flight))
        _add_labeled_paragraph(document, "航班号", flight.flight_number or "无法解析")
        _add_labeled_paragraph(document, "机型", flight.aircraft_type.value)
        _add_labeled_paragraph(
            document,
            "实际接地交叉角",
            _display_number(flight.angle_difference, "°"),
        )
        _add_labeled_paragraph(
            document,
            "着陆跑道磁航向",
            _display_number(
                None if flight.runway_heading is None else flight.runway_heading % 360,
                "°",
            ),
        )
        if flight.crosswind_metric.success and flight.crosswind_metric.value is not None:
            _add_labeled_paragraph(
                document,
                "采用的风速",
                _display_number(flight.wind_speed, " kt"),
            )
            _add_labeled_paragraph(
                document,
                "采用的风向",
                _display_number(flight.wind_direction, "°"),
            )
            _add_labeled_paragraph(
                document,
                "侧风分量",
                _display_number(flight.crosswind, " kt"),
            )
            _add_labeled_paragraph(
                document,
                "风数据来源",
                (
                    f"CSV第{flight.wind_source_row}行，"
                    f"与接地行相差{flight.wind_row_offset:+d}行"
                    if flight.wind_source_row is not None
                    and flight.wind_row_offset is not None
                    else "无法计算"
                ),
            )
        else:
            _add_labeled_paragraph(document, "采用的风速", "无法计算")
            _add_labeled_paragraph(document, "采用的风向", "无法计算")
            _add_labeled_paragraph(document, "侧风分量", "无法计算")
            _add_labeled_paragraph(document, "风数据来源", "无法计算")
        _add_labeled_paragraph(document, "源文件", flight.path.name)
        _keep_event_block_together(document, start_index)


def _add_pitch_rate_events(document: Document, flights: list[FlightResult]) -> None:
    events = [
        flight
        for flight in flights
        if flight.metrics["pitch_rate"].success
        and flight.metrics["pitch_rate"].value is True
    ]
    if not events:
        document.add_paragraph("无")
        return

    for event_index, flight in enumerate(events, start=1):
        start_index = len(document.paragraphs)
        document.add_heading(f"事件 {event_index}", level=3)
        _add_labeled_paragraph(document, "航段日期时间", _display_time(flight))
        _add_labeled_paragraph(document, "航班号", flight.flight_number or "无法解析")
        _add_labeled_paragraph(document, "机型", flight.aircraft_type.value)
        _add_labeled_paragraph(
            document,
            "最大抬头速率",
            _display_number(flight.max_pitch_rate, "°/s"),
        )
        _add_labeled_paragraph(document, "高高原起飞", _high_altitude_status(flight))
        _add_labeled_paragraph(document, "源文件", flight.path.name)
        _keep_event_block_together(document, start_index)


def _add_reverse_rudder_events(document: Document, flights: list[FlightResult]) -> None:
    events = [
        flight
        for flight in flights
        if flight.metrics["reverse_rudder"].success
        and flight.metrics["reverse_rudder"].value is True
    ]
    if not events:
        document.add_paragraph("无")
        return

    for event_index, flight in enumerate(events, start=1):
        start_index = len(document.paragraphs)
        document.add_heading(f"事件 {event_index}", level=3)
        _add_labeled_paragraph(document, "航段日期时间", _display_time(flight))
        _add_labeled_paragraph(document, "航班号", flight.flight_number or "无法解析")
        _add_labeled_paragraph(document, "机型", flight.aircraft_type.value)
        _add_labeled_paragraph(
            document,
            "接地时飞机磁航向",
            _display_number(
                None if flight.aircraft_heading is None else flight.aircraft_heading % 360,
                "°",
            ),
        )
        _add_labeled_paragraph(
            document,
            "着陆跑道磁航向",
            _display_number(
                None if flight.runway_heading is None else flight.runway_heading % 360,
                "°",
            ),
        )
        if flight.signed_angle is None:
            signed_display = "无法计算"
        elif flight.signed_angle < 0:
            signed_display = f"左交叉{abs(flight.signed_angle):.2f}°"
        else:
            signed_display = f"右交叉{abs(flight.signed_angle):.2f}°"
        _add_labeled_paragraph(document, "带方向接地交叉角", signed_display)
        _add_labeled_paragraph(
            document,
            "交叉方向",
            (
                "左交叉"
                if flight.signed_angle is not None and flight.signed_angle < 0
                else "右交叉"
            ),
        )
        if flight.reverse_rudder_value is None:
            rudder_display = "无法计算"
        elif flight.reverse_rudder_value > 0:
            rudder_display = f"左舵+{flight.reverse_rudder_value:.2f}"
        else:
            rudder_display = f"右舵{flight.reverse_rudder_value:.2f}"
        _add_labeled_paragraph(document, "代表舵量值", rudder_display)
        _add_labeled_paragraph(
            document,
            "代表舵量数据所在CSV行",
            str(flight.reverse_rudder_row or "无法计算"),
        )
        _add_labeled_paragraph(document, "50ft行", str(flight.fifty_ft_row or "无法计算"))
        _add_labeled_paragraph(document, "接地行", str(flight.touchdown_row or "无法计算"))
        _add_labeled_paragraph(document, "源CSV文件名", flight.path.name)
        _keep_event_block_together(document, start_index)


def _validate_docx(path: Path, expected_people: list[str]) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        broken_member = archive.testzip()
        if broken_member is not None:
            raise OutputError(f"DOCX压缩包校验失败：{broken_member}")

    validation_document = Document(path)
    heading_people = [
        paragraph.text
        for paragraph in validation_document.paragraphs
        if paragraph.style is not None and paragraph.style.name == "Heading 1"
    ]
    if heading_people != expected_people:
        raise OutputError("DOCX校验失败：人员标题与分析结果不一致")


def write_docx_report(output_dir: Path, result: AnalysisResult) -> Path:
    if result.cancelled:
        raise OutputError("分析已取消，不生成DOCX报告")
    output_dir.mkdir(parents=True, exist_ok=True)
    temporary_path = output_dir / f".qar_report_{uuid.uuid4().hex}.tmp.docx"

    try:
        document = Document()
        _configure_document(document)
        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("QAR批量分析报告")
        _set_font(title_run, "Microsoft YaHei", 20, bold=True)

        intro = document.add_paragraph()
        intro.paragraph_format.space_after = Pt(10)
        intro.add_run(
            "本报告汇总所选数据目录内各人员的QAR分析结果，并列出接地交叉角、抬头速率超限和疑似接地前反向蹬舵事件。"
        )
        _add_labeled_paragraph(
            document,
            "运行开始时间",
            result.started_at.strftime("%Y-%m-%d %H:%M:%S"),
        )
        _add_labeled_paragraph(
            document,
            "分析结束时间",
            result.ended_at.strftime("%Y-%m-%d %H:%M:%S"),
        )
        _add_labeled_paragraph(document, "所选数据目录", str(result.selected_folder))
        _add_labeled_paragraph(document, "人员数量", str(len(result.summaries)))

        flights_by_person: dict[str, list[FlightResult]] = {}
        for flight in result.flights:
            flights_by_person.setdefault(flight.person, []).append(flight)

        for person_index, summary in enumerate(result.summaries):
            if person_index:
                document.add_paragraph()
            document.add_heading(summary.person, level=1)
            document.add_heading("人员指标汇总", level=2)
            _add_two_column_table(document, _summary_rows(summary))
            person_flights = flights_by_person.get(summary.person, [])
            document.add_heading("接地交叉角事件", level=2)
            _add_cross_angle_events(document, person_flights)
            document.add_heading("抬头速率事件", level=2)
            _add_pitch_rate_events(document, person_flights)
            document.add_heading("疑似接地前反向蹬舵事件", level=2)
            _add_reverse_rudder_events(document, person_flights)

        document.save(temporary_path)
        expected_people = [summary.person for summary in result.summaries]
        _validate_docx(temporary_path, expected_people)
        stem = f"qar_report_{result.started_at:%Y%m%d_%H%M%S}"
        target_path = output_dir / f"{stem}.docx"
        return _publish_without_overwrite(temporary_path, target_path, "DOCX报告")
    except OutputError:
        raise
    except Exception as exc:
        raise OutputError(f"生成DOCX报告失败：{exc}") from exc
    finally:
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass
